# ============================================================
# SINH data/scenarios/*.xlsx (kịch bản keyword-driven) từ định nghĩa trong file này -> tái tạo được.
# Mỗi workbook có 2 sheet:
#   TestSteps: template_id | step | keyword | args (mảng JSON) | note
#   TestData : case_id | template_id | title | tags | requires | known_bug | <cột placeholder...>
# {cot} được gắn lúc collection; ${...} và "@data:..." được thay lúc chạy.
# Chạy:
#     .venv/bin/python scripts/gen_scenarios_xlsx.py
# ============================================================
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config import SCENARIO_DIR  # noqa: E402
from utils.hybrid_reader import DATA_COLUMNS, STEP_COLUMNS, read_hybrid_cases  # noqa: E402


def steps(template_id, *rows):
    """rows: (keyword, args, note) -> dòng TestSteps có step đánh số liên tục."""
    return [
        [template_id, str(number), keyword, json.dumps(args, ensure_ascii=False) if args else "", note]
        for number, (keyword, args, note) in enumerate(rows, 1)
    ]


WORKBOOKS = {
    "01-auth.xlsx": {
        "steps": [
            *steps(
                "T-LOGIN-ERROR",
                ("auth.openLogin", [], ""),
                ("auth.loginExpectingError", ["{username}", "{password}", "{error}"], "lỗi nguyên văn"),
                ("auth.verifyLoginFieldsHighlighted", [], ""),
                ("common.verifyUrl", ["/"], "vẫn ở trang login"),
            ),
            *steps(
                "T-PROTECTED",
                ("common.openPage", ["{path}"], "chưa đăng nhập"),
                ("auth.verifyLoginFormDisplayed", [], ""),
                ("common.verifyErrorMessage", ["{error}"], ""),
            ),
        ],
        "placeholders": ["username", "password", "error", "path"],
        "data": [
            ["KW-AUTH-01", "T-LOGIN-ERROR", "Excel: user bị khoá", "@smoke", "password", "",
             "locked_out_user", "${env.password}", "${msg.login.lockedOut}", ""],
            ["KW-AUTH-02", "T-LOGIN-ERROR", "Excel: sai mật khẩu", "", "", "",
             "standard_user", "not_the_password", "${msg.login.wrongCredentials}", ""],
            ["KW-AUTH-03", "T-PROTECTED", "Excel: mở giỏ hàng khi chưa đăng nhập", "", "", "",
             "", "", "${msg.loginRequired.cart}", "/cart.html"],
        ],
    },
    "02-shopping.xlsx": {
        "steps": [
            *steps(
                "T-BUY",
                ("auth.loginAs", ["{user}"], ""),
                ("inventory.addProductToCart", ["{product}"], ""),
                ("common.verifyCartBadge", [1], ""),
                ("cart.openCart", [], ""),
                ("cart.verifyCartItems", [["{product}"]], ""),
                ("cart.proceedToCheckout", [], ""),
                ("checkout.fillCheckoutInfo", ["Kw", "Tester", "70000"], ""),
                ("checkout.continueCheckout", [], ""),
                ("checkout.verifyItemTotal", ["{item_total}"], ""),
                ("checkout.verifyTotalEqualsItemTotalPlusTax", [], ""),
                ("checkout.finishCheckout", [], ""),
                ("checkout.verifyOrderComplete", ["${msg.checkout.completeHeader}", "${msg.checkout.completeText}"], "biến runtime"),
            ),
            *steps(
                "T-SORT",
                ("auth.loginAs", ["{user}"], ""),
                ("inventory.sortProductsBy", ["{label}"], ""),
                ("inventory.verifyActiveSort", ["{label}"], ""),
                ("inventory.verifyProductNames", ["{names}"], "@data: nạp danh sách mong đợi"),
            ),
            *steps(
                "T-RESET",
                ("auth.loginAs", ["{user}"], ""),
                ("inventory.addProductsToCart", [["{product_1}", "{product_2}"]], ""),
                ("common.verifyCartBadge", [2], ""),
                ("menu.resetAppState", [], ""),
                ("common.verifyCartBadge", [0], ""),
                ("common.verifyCartStorage", [[]], "localStorage cart-contents bị xoá"),
                ("common.reloadPage", [], "nút trên trang chỉ đổi sau reload"),
                ("inventory.verifyProductButtonState", ["{product_1}", "Add to cart"], ""),
            ),
            *steps(
                "T-CHECKOUT-REQUIRED",
                ("auth.loginAs", ["{user}"], ""),
                ("cart.openCart", [], ""),
                ("cart.proceedToCheckout", [], ""),
                ("checkout.continueCheckoutExpectingError", ["{first_name}", "{last_name}", "{postal_code}", "{error}"], ""),
                ("common.verifyUrl", ["/checkout-step-one.html"], ""),
            ),
        ],
        "placeholders": ["user", "product", "item_total", "label", "names", "product_1", "product_2",
                         "first_name", "last_name", "postal_code", "error"],
        "data": [
            ["KW-BUY-01", "T-BUY", "Excel: standard_user mua Onesie", "@smoke", "password", "",
             "standard_user", "Sauce Labs Onesie", "$7.99", "", "", "", "", "", "", "", ""],
            ["KW-BUY-02", "T-BUY", "Excel: performance_glitch_user mua Fleece Jacket", "", "password", "",
             "performance_glitch_user", "Sauce Labs Fleece Jacket", "$49.99", "", "", "", "", "", "", "", ""],
            ["KW-BUY-03", "T-BUY", "Excel: error_user hoàn tất đơn", "", "password",
             "error_user bấm Finish không có tác dụng, đứng ở checkout-step-two",
             "error_user", "Sauce Labs Backpack", "$29.99", "", "", "", "", "", "", "", ""],
            ["KW-SORT-01", "T-SORT", "Excel: sắp xếp giá cao -> thấp", "", "password", "",
             "standard_user", "", "", "Price (high to low)", "@data:inventory/sorting.json#hilo.names", "", "", "", "", "", ""],
            ["KW-SORT-02", "T-SORT", "Excel: problem_user sắp xếp Z -> A", "", "password",
             "problem_user chọn kiểu sắp xếp nhưng thứ tự sản phẩm không đổi",
             "problem_user", "", "", "Name (Z to A)", "@data:inventory/sorting.json#za.names", "", "", "", "", "", ""],
            ["KW-RESET-01", "T-RESET", "Excel: Reset App State xoá giỏ", "", "password", "",
             "standard_user", "", "", "", "", "Sauce Labs Backpack", "Sauce Labs Bike Light", "", "", "", ""],
            ["KW-CHK-01", "T-CHECKOUT-REQUIRED", "Excel: checkout thiếu Postal Code", "", "password", "",
             "standard_user", "", "", "", "", "", "", "Hai", "Nguyen", "", "${msg.checkout.postalCodeRequired}"],
        ],
    },
}


def write_workbook(path, spec):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TestSteps"
    sheet.append(STEP_COLUMNS)
    for row in spec["steps"]:
        sheet.append(row)
    data = workbook.create_sheet("TestData")
    data.append(DATA_COLUMNS + spec["placeholders"])
    for row in spec["data"]:
        if len(row) != len(DATA_COLUMNS) + len(spec["placeholders"]):
            raise ValueError(f"{path.name} {row[0]}: số cột không khớp header")
        data.append(row)
    for ws in (sheet, data):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        # Giữ chuỗi như "$7.99", "70000" không bị Excel đổi kiểu.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.number_format = "@"
    workbook.save(path)


def main():
    SCENARIO_DIR.mkdir(parents=True, exist_ok=True)
    for name, spec in WORKBOOKS.items():
        path = SCENARIO_DIR / name
        write_workbook(path, spec)
        cases = read_hybrid_cases(path)  # tự kiểm tra: đọc lại qua đúng reader của framework
        print(f"{name}: {len(cases)} kịch bản")
    return 0


if __name__ == "__main__":
    sys.exit(main())
